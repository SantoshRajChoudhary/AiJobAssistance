"""
Shared pytest fixtures for the AI Job Assistant test suite.
"""
import uuid
import pytest
import fitz
import io
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.storage_service import storage_service


@pytest.fixture(scope="session")
def client():
    """Reusable FastAPI TestClient (session-scoped for performance)."""
    return TestClient(app)


@pytest.fixture
def unique_session():
    """Provides a guaranteed-unique session ID for each test."""
    return f"test_sess_{uuid.uuid4().hex[:10]}"


@pytest.fixture(scope="session")
def sample_jobs():
    """Returns the first 5 jobs from the dataset for use in tests."""
    jobs = storage_service.load_jobs()
    assert len(jobs) >= 5, "Dataset must have at least 5 jobs for tests"
    return jobs[:5]


@pytest.fixture
def sample_pdf_bytes():
    """Creates an in-memory PDF with a realistic resume."""
    text = (
        "Alice Dev\n"
        "alice@example.com\n"
        "+91-9876543210\n"
        "Software Engineer\n"
        "Skills: Python, FastAPI, Docker, PostgreSQL, React\n"
        "WORK EXPERIENCE\n"
        "Backend Developer at CloudCo\n"
        "Jan 2022 - Present\n"
        "Built microservices and REST APIs\n"
        "EDUCATION\n"
        "B.Tech CSE, NIT Trichy, 2022\n"
        "PROJECTS\n"
        "AI Job Assistant - Python, React, FastAPI\n"
    )
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes():
    """Creates an in-memory DOCX with a realistic resume."""
    text = (
        "Bob Smith\n"
        "bob@example.org\n"
        "Frontend Engineer\n"
        "Skills: React, TypeScript, Next.js, Tailwind CSS\n"
        "WORK EXPERIENCE\n"
        "Frontend Dev at Nova Tech\n"
        "EDUCATION\n"
        "B.Tech in Computer Science, 2023\n"
    )
    doc = Document()
    doc.add_heading("Resume", level=1)
    for line in text.splitlines():
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def backend_resume():
    """A ready-to-use backend engineer resume dict for matcher tests."""
    return {
        "name": "Test User",
        "email": "test@example.com",
        "designation": "Backend Engineer",
        "skills": ["Python", "FastAPI", "Docker", "PostgreSQL", "AWS"],
        "total_experience_years": 2.0,
        "contact": {"location": "Remote"},
    }


@pytest.fixture
def frontend_resume():
    """A ready-to-use frontend engineer resume dict for matcher tests."""
    return {
        "name": "Frontend User",
        "email": "fe@example.com",
        "designation": "Frontend Engineer",
        "skills": ["React", "TypeScript", "Next.js", "Tailwind CSS"],
        "total_experience_years": 1.0,
        "contact": {"location": "Bangalore"},
    }
