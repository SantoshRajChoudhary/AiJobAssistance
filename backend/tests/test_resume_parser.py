"""
Tests for Resume Parser Tool.
Covers:
  - Valid PDF parsing (using PyMuPDF)
  - Valid DOCX parsing (using python-docx)
  - Empty or 0-byte files
  - Unreadable / corrupted files
  - Missing resume fields defaulting to schema defaults (null or empty lists)
"""

import io
import pytest
import fitz  # PyMuPDF
from docx import Document

from app.agents.tools.resume_parser import (
    extract_text_from_pdf,
    extract_text_from_docx,
    parse_resume_document,
)
from app.utils.file_validation import validate_resume_file
from app.utils.errors import InvalidInputError, FileTooLargeError


def create_sample_pdf(text: str) -> bytes:
    """Creates an in-memory PDF with given text using PyMuPDF."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), text)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def create_sample_docx(text: str) -> bytes:
    """Creates an in-memory DOCX with given text using python-docx."""
    doc = Document()
    doc.add_heading("Resume", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_valid_pdf_parsing():
    sample_text = (
        "Jane Doe\n"
        "jane.doe@example.com\n"
        "+1-555-123-4567\n"
        "Software Engineer\n"
        "Skills: Python, FastAPI, Docker, PostgreSQL\n"
        "WORK EXPERIENCE\n"
        "Senior Backend Developer at Acme Corp\n"
        "Jan 2021 - Present\n"
        "Built resilient microservices and APIs\n"
        "EDUCATION\n"
        "B.S. Computer Science, University of Technology\n"
        "PROJECTS\n"
        "API Gateway - Python, Docker\n"
    )
    pdf_bytes = create_sample_pdf(sample_text)
    extracted = extract_text_from_pdf(pdf_bytes)
    assert "jane.doe@example.com" in extracted
    assert "Python" in extracted

    profile = await parse_resume_document(pdf_bytes, ".pdf")
    assert profile.email == "jane.doe@example.com"
    assert "Python" in profile.skills or "FastAPI" in profile.skills
    assert profile.name is not None


@pytest.mark.asyncio
async def test_valid_docx_parsing():
    sample_text = (
        "John Smith\n"
        "john.smith@example.org\n"
        "Frontend Engineer\n"
        "Skills: React, TypeScript, Tailwind CSS\n"
        "WORK EXPERIENCE\n"
        "Frontend Dev at Nova Tech\n"
        "EDUCATION\n"
        "B.Tech in Computer Science\n"
    )
    docx_bytes = create_sample_docx(sample_text)
    extracted = extract_text_from_docx(docx_bytes)
    assert "john.smith@example.org" in extracted
    assert "React" in extracted

    profile = await parse_resume_document(docx_bytes, ".docx")
    assert profile.email == "john.smith@example.org"
    assert "React" in profile.skills or "TypeScript" in profile.skills


def test_empty_file_validation():
    with pytest.raises(InvalidInputError) as exc_info:
        validate_resume_file("empty.pdf", "application/pdf", b"")
    assert "empty" in str(exc_info.value).lower()


def test_file_too_large_validation():
    # 6 MB dummy payload
    oversized = b"x" * (6 * 1024 * 1024)
    with pytest.raises(FileTooLargeError):
        validate_resume_file("large.pdf", "application/pdf", oversized)


def test_invalid_file_extension():
    with pytest.raises(InvalidInputError) as exc_info:
        validate_resume_file("script.exe", "application/octet-stream", b"sample")
    assert "unsupported" in str(exc_info.value).lower()


@pytest.mark.asyncio
async def test_corrupted_pdf_handling():
    corrupted_bytes = b"%PDF-1.4 definitely not a valid pdf header content"
    with pytest.raises(InvalidInputError):
        extract_text_from_pdf(corrupted_bytes)


@pytest.mark.asyncio
async def test_missing_fields_default_to_null_and_empty_arrays():
    """Missing fields must strictly default to null or empty lists, never fabricated."""
    minimal_text = "Developer\nSkills: Python\n"
    pdf_bytes = create_sample_pdf(minimal_text)
    profile = await parse_resume_document(pdf_bytes, ".pdf")

    # Name, phone, etc. might be absent
    assert profile.phone is None or profile.phone == ""
    assert isinstance(profile.skills, list)
    assert isinstance(profile.experience, list)
    assert isinstance(profile.education, list)
    assert isinstance(profile.projects, list)
