"""
Resume Parser Tool.
Extracts raw text from PDF (using PyMuPDF) and DOCX (using python-docx),
detects unreadable/empty files, and structures into a validated ResumeProfile schema.
"""

import io
import fitz  # PyMuPDF
from docx import Document
from typing import Dict, Any

from app.utils.errors import InvalidInputError
from app.models.schemas import ResumeProfile
from app.services.llm_service import llm_service


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF using PyMuPDF (fitz)."""
    text_parts = []
    try:
        # Open PDF from byte stream
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text()
            if text and text.strip():
                text_parts.append(text.strip())
        doc.close()
    except Exception as e:
        raise InvalidInputError(f"Failed to read PDF file. File may be corrupted or encrypted: {e}")

    extracted = "\n".join(text_parts).strip()
    if not extracted:
        raise InvalidInputError("The uploaded PDF is empty or contains only non-extractable scanned images.")
    return extracted


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX using python-docx."""
    text_parts = []
    try:
        doc = Document(io.BytesIO(file_bytes))
        for p in doc.paragraphs:
            if p.text and p.text.strip():
                text_parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        text_parts.append(cell.text.strip())
    except Exception as e:
        raise InvalidInputError(f"Failed to read DOCX file. File may be corrupted: {e}")

    extracted = "\n".join(text_parts).strip()
    if not extracted:
        raise InvalidInputError("The uploaded Word document is empty or contains no extractable text.")
    return extracted


async def parse_resume_document(file_bytes: bytes, extension: str) -> ResumeProfile:
    """
    Parses resume document bytes into a validated ResumeProfile Pydantic object.
    Uses PyMuPDF for PDF and python-docx for DOCX, then structures with LLMService.
    """
    if extension == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
    elif extension in {".docx", ".doc"}:
        raw_text = extract_text_from_docx(file_bytes)
    else:
        raise InvalidInputError(f"Unsupported file format '{extension}'. Use PDF or DOCX.")

    # Structure text into validated JSON
    structured_data = await llm_service.parse_resume_to_json(raw_text)

    # Validate against Pydantic schema
    try:
        profile = ResumeProfile(**structured_data)
        return profile
    except Exception as e:
        # Fallback to minimal valid profile if schema validation fails
        return ResumeProfile(
            name=structured_data.get("name"),
            email=structured_data.get("email"),
            phone=structured_data.get("phone"),
            designation=structured_data.get("designation"),
            skills=structured_data.get("skills", []),
        )
